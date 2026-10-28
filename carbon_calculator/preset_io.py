# SPDX-License-Identifier: MIT
# -*- coding: utf-8 -*-
"""프리셋 엑셀(.xlsx) — 지역·수종 인벤토리를 한 번에 불러오고 양식을 만든다.

수종마다 [+ 추가] 대화상자를 거치지 않고, 정해진 양식의 엑셀을 불러와 지역 탭을
만들고 바로 계산할 수 있게 한다. Qt 에 의존하지 않는다(테스트 용이).

양식 (시트 이름·머리글은 국문/영문 모두 인식)::

    [지역]  지역명 | 가로(m) | 세로(m) | 대상지 유형 | 위도 | 경도
    [수종]  지역명 | 구분(교목/관목) | 수종 | 직경(cm) | 개체수
            | 현재 생존율(%) | 연간 생존율(%) | 두 번째 변수

- 직경은 교목 DBH, 관목 RCD 이며 모두 cm.
- 생존율 두 열은 비우면 100 % (전량 생존).
- [지역] 시트가 없으면 [수종] 시트의 지역명으로 20 m × 20 m 지역을 만든다.
- 위도·경도(십진수 도, WGS84)는 선택 입력이다. 둘 다 적어야 위치로 쓴다.
- 수종은 국명(라이브러리 표기 그대로) 또는 학명으로 적을 수 있다.
"""
from __future__ import annotations

import math
import re
from dataclasses import dataclass, field
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from . import i18n
from .i18n import tr
from .survival import DEFAULT_SURVIVAL_PCT

KIND_TREE, KIND_SHRUB = "tree", "shrub"
DEFAULT_SIDE_M = 20

# ── 머리글 인식용 별칭 (공백·대소문자·괄호 단위 무시) ──
_REGION_SHEET = ("지역", "regions", "region", "sites", "site")
_SPECIES_SHEET = ("수종", "species", "inventory", "인벤토리")
_COLUMNS = {
    "region": ("지역명", "지역", "region", "site", "sitename", "regionname", "profile"),
    "width": ("가로", "가로m", "width", "widthm"),
    "height": ("세로", "세로m", "height", "heightm", "length"),
    "env": ("대상지유형", "대상지", "유형", "sitecategory", "category", "environment"),
    "lat": ("위도", "latitude", "lat"),
    "lon": ("경도", "longitude", "lon", "lng", "long"),
    "kind": ("구분", "생활형", "type", "kind", "growthform"),
    "species": ("수종", "수종명", "species", "speciesname"),
    "diameter": ("직경", "직경cm", "dbh/rcd", "dbhrcd", "dbh", "rcd", "diameter", "diametercm"),
    "quantity": ("개체수", "수량", "본수", "quantity", "count", "n", "individuals"),
    "survival_initial": ("현재생존율", "생존율", "현재생존률", "생존률", "initialsurvival", "survival"),
    "survival_annual": ("연간생존율", "연간생존률", "annualsurvival", "annualsurvivalrate"),
    "var2": ("두번째변수", "변수2", "secondvariable", "var2"),
}
_KIND_WORDS = {
    KIND_TREE: ("교목", "tree", "trees", "t"),
    KIND_SHRUB: ("관목", "shrub", "shrubs", "s"),
}


@dataclass
class PresetRow:
    kind: str
    species: str                      # 라이브러리 내부 키(국명)
    diameter: float
    quantity: int
    survival_initial: float = DEFAULT_SURVIVAL_PCT
    survival_annual: float = DEFAULT_SURVIVAL_PCT
    var2: float | None = None
    source_row: int = 0               # 엑셀 행 번호 (오류 안내용)


@dataclass
class PresetRegion:
    name: str
    width_m: int = DEFAULT_SIDE_M
    height_m: int = DEFAULT_SIDE_M
    environment: str = ""
    rows: list[PresetRow] = field(default_factory=list)
    latitude: float | None = None     # 십진수 도 (WGS84), 선택
    longitude: float | None = None


def validate_location(lat, lon) -> tuple[float, float] | None:
    """위도·경도 쌍을 검사한다. 둘 다 비면 None, 하나만 있거나 범위 밖이면 ValueError."""
    blank = [v is None or (isinstance(v, str) and not v.strip()) for v in (lat, lon)]
    if all(blank):
        return None
    if any(blank):
        raise ValueError(tr("위도와 경도를 함께 입력해 주세요"))
    try:
        lat_v, lon_v = (float(str(v).replace(",", "").strip()) if isinstance(v, str) else float(v)
                        for v in (lat, lon))
    except ValueError:
        raise ValueError(tr("위도·경도는 숫자(십진수 도)로 입력해 주세요"))
    if not (math.isfinite(lat_v) and -90 <= lat_v <= 90):
        raise ValueError(tr("위도는 -90 ~ 90 사이여야 합니다"))
    if not (math.isfinite(lon_v) and -180 <= lon_v <= 180):
        raise ValueError(tr("경도는 -180 ~ 180 사이여야 합니다"))
    return lat_v, lon_v


def format_location(lat: float, lon: float) -> str:
    """36.4986, 127.2896 → '36.4986°N, 127.2896°E'."""
    ns = "N" if lat >= 0 else "S"
    ew = "E" if lon >= 0 else "W"
    return f"{abs(lat):.4f}°{ns}, {abs(lon):.4f}°{ew}"


@dataclass
class PresetResult:
    regions: list[PresetRegion]
    errors: list[str]                 # 해당 행을 건너뛴 문제
    notes: list[str]                  # 진행에는 문제없는 안내


def _norm(text) -> str:
    """머리글 비교용: 소문자, 공백·괄호 속 단위·기호 제거."""
    s = str(text or "").strip().lower()
    s = re.sub(r"\((m|cm|%|%/년|%/yr|%/year)\)", "", s)
    return re.sub(r"[\s_()\[\]%·\-]", "", s)


def _match_columns(header: Iterable) -> dict[str, int]:
    found: dict[str, int] = {}
    for idx, cell in enumerate(header):
        key = _norm(cell)
        if not key:
            continue
        for name, aliases in _COLUMNS.items():
            if name in found:
                continue
            if key in {_norm(a) for a in aliases}:
                found[name] = idx
                break
    return found


def _find_sheet(wb, names: tuple, required: tuple[str, ...]):
    """이름이 맞는 시트 → 없으면 필요한 머리글을 가진 첫 시트."""
    wanted = {_norm(n) for n in names}
    for ws in wb.worksheets:
        if _norm(ws.title) in wanted:
            return ws
    for ws in wb.worksheets:
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        cols = _match_columns(header)
        if all(r in cols for r in required):
            return ws
    return None


def _number(value, what: str, row: int, *, allow_blank=False, default=None) -> float | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        if allow_blank:
            return default
        raise ValueError(tr("{row}행: {what} 값이 비어 있습니다").format(row=row, what=what))
    try:
        v = float(str(value).replace(",", "").strip()) if isinstance(value, str) else float(value)
    except ValueError:
        raise ValueError(tr("{row}행: {what} 값이 숫자가 아닙니다 ({value})")
                         .format(row=row, what=what, value=value))
    if not math.isfinite(v):
        raise ValueError(tr("{row}행: {what} 값이 숫자가 아닙니다 ({value})")
                         .format(row=row, what=what, value=value))
    return v


def species_aliases(names: Iterable[str]) -> dict[str, str]:
    """국명·학명(꼬리표 포함) → 라이브러리 키. 학명이 겹치면 그 학명 단독 별칭은 뺀다."""
    aliases: dict[str, str] = {}
    latin_owner: dict[str, list[str]] = {}
    for key in names:
        aliases[_norm(key)] = key
        m = re.match(r"^\s*(?P<base>[^(]+?)\s*(\((?P<qual>.*)\))?\s*$", key)
        base = m.group("base") if m else key
        latin = i18n.scientific_name(base)
        if not latin:
            continue
        latin_owner.setdefault(_norm(latin), []).append(key)
        if m and m.group("qual"):
            qual = i18n._translate_qualifier(m.group("qual"))
            aliases.setdefault(_norm(f"{latin} ({qual})"), key)
    for latin, owners in latin_owner.items():
        if len(owners) == 1:
            aliases.setdefault(latin, owners[0])
    return aliases


def _parse_kind(value) -> str | None:
    k = _norm(value)
    for kind, words in _KIND_WORDS.items():
        if k in {_norm(w) for w in words}:
            return kind
    return None


def _environment_lookup(environments: Iterable[str]) -> dict[str, str]:
    out = {}
    for env in environments:
        out[_norm(env)] = env
        out[_norm(i18n.ENVIRONMENT_EN.get(env, env))] = env
    return out


def read_preset(path: str, tree_names: Iterable[str], shrub_names: Iterable[str],
                environments: Iterable[str]) -> PresetResult:
    """프리셋 엑셀을 읽어 지역 목록을 만든다. 잘못된 행은 건너뛰고 errors 에 남긴다."""
    tree_names, shrub_names = list(tree_names), list(shrub_names)
    environments = list(environments)
    default_env = environments[0] if environments else ""
    env_lookup = _environment_lookup(environments)
    tree_alias = species_aliases(tree_names)
    shrub_alias = species_aliases(shrub_names)

    wb = load_workbook(path, data_only=True, read_only=True)
    errors: list[str] = []
    notes: list[str] = []
    try:
        sp_ws = _find_sheet(wb, _SPECIES_SHEET, ("species", "diameter", "quantity"))
        if sp_ws is None:
            raise ValueError(tr("'수종' 시트(수종·직경·개체수 열)를 찾을 수 없습니다."))
        rg_ws = _find_sheet(wb, _REGION_SHEET, ("region", "width", "height"))
        if rg_ws is sp_ws:
            rg_ws = None

        regions: dict[str, PresetRegion] = {}
        if rg_ws is not None:
            rows = rg_ws.iter_rows(values_only=True)
            cols = _match_columns(next(rows, ()))
            for r_idx, values in enumerate(rows, start=2):
                get = lambda c: values[cols[c]] if c in cols and cols[c] < len(values) else None  # noqa: E731
                name = str(get("region") or "").strip()
                if not name:
                    continue
                try:
                    w = _number(get("width"), tr("가로(m)"), r_idx, allow_blank=True, default=DEFAULT_SIDE_M)
                    h = _number(get("height"), tr("세로(m)"), r_idx, allow_blank=True, default=DEFAULT_SIDE_M)
                    if w <= 0 or h <= 0:
                        raise ValueError(tr("{row}행: 가로·세로는 0보다 커야 합니다").format(row=r_idx))
                except ValueError as exc:
                    errors.append(f"[{rg_ws.title}] {exc}")
                    continue
                env_raw = get("env")
                env = env_lookup.get(_norm(env_raw), None) if env_raw else default_env
                if env is None:
                    notes.append(tr("[{sheet}] {row}행: 대상지 유형 '{value}' 을(를) 알 수 없어 '{default}' 로 둡니다")
                                 .format(sheet=rg_ws.title, row=r_idx, value=env_raw, default=default_env))
                    env = default_env
                if name in regions:
                    notes.append(tr("[{sheet}] {row}행: 지역명 '{name}' 이(가) 중복되어 앞의 값을 씁니다")
                                 .format(sheet=rg_ws.title, row=r_idx, name=name))
                    continue
                try:
                    location = validate_location(get("lat"), get("lon"))
                except ValueError as exc:
                    notes.append(tr("[{sheet}] {row}행: {error} — 위치 없이 불러옵니다")
                                 .format(sheet=rg_ws.title, row=r_idx, error=exc))
                    location = None
                regions[name] = PresetRegion(name, int(round(w)), int(round(h)), env,
                                             latitude=location[0] if location else None,
                                             longitude=location[1] if location else None)

        rows = sp_ws.iter_rows(values_only=True)
        cols = _match_columns(next(rows, ()))
        missing = [c for c in ("species", "diameter", "quantity") if c not in cols]
        if missing:
            raise ValueError(tr("'{sheet}' 시트에 필요한 열이 없습니다: {cols}")
                             .format(sheet=sp_ws.title, cols=", ".join(missing)))
        single_name = next(iter(regions)) if len(regions) == 1 else None
        for r_idx, values in enumerate(rows, start=2):
            get = lambda c: values[cols[c]] if c in cols and cols[c] < len(values) else None  # noqa: E731
            raw_species = get("species")
            if raw_species is None or not str(raw_species).strip():
                continue
            where = f"[{sp_ws.title}] "
            region_name = str(get("region") or "").strip() or single_name or tr("불러온 데이터")
            kind = _parse_kind(get("kind")) if get("kind") is not None else None
            key_norm = _norm(raw_species)
            if kind is None:
                kind = (KIND_TREE if key_norm in tree_alias
                        else KIND_SHRUB if key_norm in shrub_alias else None)
            alias = tree_alias if kind == KIND_TREE else shrub_alias
            species = alias.get(key_norm) if kind else None
            if species is None:
                errors.append(where + tr("{row}행: '{species}' 은(는) {kind} 라이브러리에 없는 수종입니다")
                              .format(row=r_idx, species=raw_species,
                                      kind=tr("교목") if kind == KIND_TREE else tr("관목") if kind else tr("교목·관목")))
                continue
            try:
                d = _number(get("diameter"), tr("직경(cm)"), r_idx)
                q = _number(get("quantity"), tr("개체수"), r_idx)
                s0 = _number(get("survival_initial"), tr("현재 생존율(%)"), r_idx,
                             allow_blank=True, default=DEFAULT_SURVIVAL_PCT)
                s1 = _number(get("survival_annual"), tr("연간 생존율(%)"), r_idx,
                             allow_blank=True, default=DEFAULT_SURVIVAL_PCT)
                v2 = _number(get("var2"), tr("두 번째 변수"), r_idx, allow_blank=True, default=None)
                if d <= 0:
                    raise ValueError(tr("{row}행: 직경은 0보다 커야 합니다").format(row=r_idx))
                if q < 0 or abs(q - round(q)) > 1e-9:
                    raise ValueError(tr("{row}행: 개체수는 0 이상의 정수여야 합니다").format(row=r_idx))
                for v in (s0, s1):
                    if not 0 <= v <= 100:
                        raise ValueError(tr("{row}행: 생존율은 0~100 % 사이여야 합니다").format(row=r_idx))
            except ValueError as exc:
                errors.append(where + str(exc))
                continue
            region = regions.get(region_name)
            if region is None:
                region = regions[region_name] = PresetRegion(region_name, environment=default_env)
                if rg_ws is not None:
                    notes.append(tr("지역 '{name}' 이(가) 지역 시트에 없어 {w} m × {h} m 로 만듭니다")
                                 .format(name=region_name, w=DEFAULT_SIDE_M, h=DEFAULT_SIDE_M))
            region.rows.append(PresetRow(kind, species, d, int(round(q)), s0, s1, v2, r_idx))
    finally:
        wb.close()

    result = [r for r in regions.values() if r.rows]
    empty = [r.name for r in regions.values() if not r.rows]
    if empty:
        notes.append(tr("수종 입력이 없어 건너뛴 지역: {names}").format(names=", ".join(empty)))
    return PresetResult(result, errors, notes)


# ─────────────────────────── 양식 쓰기 ───────────────────────────

_HEADER_FILL = PatternFill("solid", fgColor="2E8B57")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_NOTE_FONT = Font(color="555555")


# 예시 좌표 — 입력 방법을 보여 주기 위한 값(국립세종수목원 부근)이며 Profile 의 실제 위치가 아니다.
EXAMPLE_LOCATION = (36.4986, 127.2896)


def example_regions() -> list[PresetRegion]:
    """최종보고서 표3의 Profile 1~3 (동일 면적 400 m²) 인벤토리."""
    def rows(items):
        return [PresetRow(k, s, d, n) for k, s, d, n in items]
    return [
        PresetRegion("Profile 1", 20, 20, "", rows([
            (KIND_TREE, "소나무", 5.0, 30), (KIND_TREE, "졸참나무", 4.0, 20),
            (KIND_TREE, "신갈나무", 8.0, 15),
            (KIND_SHRUB, "산철쭉", 1.2, 40), (KIND_SHRUB, "병꽃나무", 1.5, 25)]),
            *EXAMPLE_LOCATION),
        PresetRegion("Profile 2", 20, 20, "", rows([
            (KIND_TREE, "소나무", 7.0, 60), (KIND_TREE, "졸참나무", 8.0, 25),
            (KIND_TREE, "신갈나무", 10.0, 15), (KIND_SHRUB, "병꽃나무", 1.5, 25)])),
        PresetRegion("Profile 3", 20, 20, "", rows([
            (KIND_TREE, "졸참나무", 4.0, 20), (KIND_TREE, "신갈나무", 6.0, 10),
            (KIND_SHRUB, "산철쭉", 1.8, 120), (KIND_SHRUB, "병꽃나무", 2.0, 90)])),
    ]


def _header(ws, labels: list[str], widths: list[int]) -> None:
    for c, (text, w) in enumerate(zip(labels, widths), 1):
        cell = ws.cell(1, c, text)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "A2"


def write_template(path: str, environments: list[str], library: dict[str, list[tuple]],
                   regions: list[PresetRegion] | None = None) -> None:
    """빈칸 없이 바로 불러올 수 있는 예시가 든 양식을 저장한다.

    library: {"tree": [(key, predictor_label, range_text), ...], "shrub": [...]}
    """
    regions = example_regions() if regions is None else regions
    default_env = environments[0] if environments else ""
    tree_word, shrub_word = tr("교목"), tr("관목")
    wb = Workbook()

    ws = wb.active
    ws.title = tr("지역")
    _header(ws, [tr("지역명"), tr("가로(m)"), tr("세로(m)"), tr("대상지 유형"),
                 tr("위도"), tr("경도")], [24, 10, 10, 24, 12, 12])
    for i, reg in enumerate(regions, start=2):
        ws.cell(i, 1, reg.name)
        ws.cell(i, 2, reg.width_m)
        ws.cell(i, 3, reg.height_m)
        ws.cell(i, 4, i18n.environment_name(reg.environment or default_env))
        if reg.latitude is not None and reg.longitude is not None:
            ws.cell(i, 5, reg.latitude).number_format = "0.000000"
            ws.cell(i, 6, reg.longitude).number_format = "0.000000"
    env_names = [i18n.environment_name(e) for e in environments]
    if env_names:
        dv = DataValidation(type="list", formula1='"' + ",".join(env_names) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add("D2:D500")

    sp = wb.create_sheet(tr("수종"))
    _header(sp, [tr("지역명"), tr("구분"), tr("수종"), tr("직경(cm)"), tr("개체수"),
                 tr("현재 생존율(%)"), tr("연간 생존율(%)"), tr("두 번째 변수")],
            [24, 8, 26, 10, 9, 14, 14, 13])
    r = 2
    for reg in regions:
        for row in reg.rows:
            sp.cell(r, 1, reg.name)
            sp.cell(r, 2, tree_word if row.kind == KIND_TREE else shrub_word)
            sp.cell(r, 3, i18n.species_name(row.species))
            sp.cell(r, 4, row.diameter)
            sp.cell(r, 5, row.quantity)
            sp.cell(r, 6, row.survival_initial)
            sp.cell(r, 7, row.survival_annual)
            if row.var2 is not None:
                sp.cell(r, 8, row.var2)
            r += 1
    dv_kind = DataValidation(type="list", formula1=f'"{tree_word},{shrub_word}"', allow_blank=True)
    sp.add_data_validation(dv_kind)
    dv_kind.add("B2:B2000")

    lib = wb.create_sheet(tr("수종목록"))
    _header(lib, [tr("구분"), tr("수종"), tr("설명변수"), tr("적합 직경범위")], [8, 30, 22, 24])
    r = 2
    for kind, word in ((KIND_TREE, tree_word), (KIND_SHRUB, shrub_word)):
        for key, predictor, rng in library.get(kind, []):
            lib.cell(r, 1, word)
            lib.cell(r, 2, i18n.species_name(key))
            lib.cell(r, 3, predictor)
            lib.cell(r, 4, rng)
            r += 1

    guide = wb.create_sheet(tr("안내"))
    guide.column_dimensions["A"].width = 100
    lines = [
        tr("ForeCarb 입력 데이터 양식"),
        "",
        tr("1. [지역] 시트: 지역명, 가로·세로(m), 대상지 유형을 적습니다. 지역마다 탭이 하나씩 만들어집니다."),
        tr("   위도·경도는 선택 입력입니다 (십진수 도, 예: 36.4986 / 127.2896). 둘 다 적으면 지역 위치로 표시됩니다."),
        tr("2. [수종] 시트: 한 줄에 한 수종(같은 직경·생존율 묶음)을 적습니다. 지역명은 [지역] 시트와 같아야 합니다."),
        tr("3. 구분은 교목 또는 관목, 수종은 [수종목록] 시트의 이름(국명 또는 학명)을 그대로 씁니다."),
        tr("4. 직경은 cm 단위입니다 (교목 DBH, 관목 RCD). 적합 직경범위 밖의 행은 계산에서 제외됩니다."),
        tr("5. 현재 생존율(%)은 0년차 수량에, 연간 생존율(%)은 해마다 누적 적용됩니다: N(t) = N × 현재 생존율 × 연간 생존율^t"),
        tr("   비워 두면 100 %(전량 생존)이며, 기존 결과(전량 생존)와 생존률 반영 결과가 함께 표시됩니다."),
        tr("6. 두 번째 변수는 수고·임분밀도 등 보조변수가 필요한 식에만 적습니다 (비우면 기본값)."),
        tr("7. ForeCarb 메뉴 [파일] → [데이터 불러오기] 로 불러오면 지역 탭이 만들어지고 바로 계산됩니다."),
        tr("8. [파일] → [데이터 저장] 은 지금 열린 지역들의 입력을 이 양식으로 저장합니다."),
        "",
        tr("예시 데이터는 최종보고서 표3의 Profile 1~3 (400 m²) 입니다."),
        tr("Profile 1 의 위도·경도는 입력 방법을 보여 주기 위한 예시 좌표(국립세종수목원 부근)입니다."),
    ]
    for i, text in enumerate(lines, start=1):
        cell = guide.cell(i, 1, text)
        cell.font = Font(bold=True, size=13, color="246B43") if i == 1 else _NOTE_FONT
    wb.save(path)
