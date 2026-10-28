# SPDX-License-Identifier: MIT
"""Survival-adjusted projections and preset workbook import/export."""

from __future__ import annotations

from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

import numpy as np
from openpyxl import Workbook, load_workbook

from carbon_calculator import i18n, preset_io, species_library as lib
from carbon_calculator.data import RESTORATION_ENVIRONMENTS
from carbon_calculator.excel_export import export_all_regions_to_excel, export_carbon1_to_excel
from carbon_calculator.i18n import tr
from carbon_calculator.projections import project_record
from carbon_calculator.survival import is_default, surviving_individuals, survival_factors

ENV = RESTORATION_ENVIRONMENTS[0]


def _names(kind):
    return list(lib.records_for_kind(kind, ENV).keys())


def _library():
    return {kind: [(n, r.var1_label, r.range_text()) for n, r in lib.records_for_kind(kind, ENV).items()]
            for kind in (lib.KIND_TREE, lib.KIND_SHRUB)}


class SurvivalTests(unittest.TestCase):
    def setUp(self) -> None:
        i18n.set_language(i18n.LANG_KO)

    def test_factors_follow_initial_times_annual_power(self) -> None:
        years = np.arange(31)
        f = survival_factors(90.0, 97.0, years)
        self.assertAlmostEqual(f[0], 0.90)
        self.assertAlmostEqual(f[10], 0.90 * 0.97 ** 10)
        self.assertAlmostEqual(f[30], 0.90 * 0.97 ** 30)
        n = surviving_individuals(30, 90.0, 97.0, years)
        self.assertAlmostEqual(n[30], 30 * 0.90 * 0.97 ** 30)

    def test_default_survival_keeps_the_all_survive_projection(self) -> None:
        rec = lib.records_for_kind(lib.KIND_TREE, ENV)["소나무"]
        years, _p, curve = project_record(rec, 5.0, 30, None)
        self.assertTrue(is_default(100, 100))
        np.testing.assert_allclose(curve * survival_factors(100, 100, years), curve)

    def test_survival_scales_carbon_linearly(self) -> None:
        rec = lib.records_for_kind(lib.KIND_TREE, ENV)["소나무"]
        years, _p, curve = project_record(rec, 5.0, 30, None)
        adjusted = curve * survival_factors(80.0, 95.0, years)
        # 생존 개체 수만큼 다시 계산한 값과 같다 (탄소량 ∝ 개체 수).
        per_tree = curve / 30
        np.testing.assert_allclose(adjusted, per_tree * 30 * 0.8 * 0.95 ** years)

    def test_out_of_range_percentages_are_rejected(self) -> None:
        for bad in (-1, 100.5, float("nan")):
            with self.subTest(value=bad), self.assertRaises(ValueError):
                survival_factors(bad, 100, [0])


class PresetWorkbookTests(unittest.TestCase):
    def setUp(self) -> None:
        i18n.set_language(i18n.LANG_KO)

    def test_template_round_trip_reproduces_report_profiles(self) -> None:
        with TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "preset.xlsx")
            preset_io.write_template(path, list(RESTORATION_ENVIRONMENTS), _library())
            result = preset_io.read_preset(path, _names(lib.KIND_TREE), _names(lib.KIND_SHRUB),
                                           RESTORATION_ENVIRONMENTS)
        self.assertEqual(result.errors, [])
        self.assertEqual([r.name for r in result.regions], ["Profile 1", "Profile 2", "Profile 3"])
        self.assertEqual([len(r.rows) for r in result.regions], [5, 4, 4])
        p1 = result.regions[0]
        self.assertEqual((p1.width_m, p1.height_m, p1.environment), (20, 20, ENV))
        first = p1.rows[0]
        self.assertEqual((first.kind, first.species, first.diameter, first.quantity),
                         ("tree", "소나무", 5.0, 30))
        self.assertEqual((first.survival_initial, first.survival_annual), (100.0, 100.0))
        # Profile 1 탄소저장량은 보고서 값(198.89 kgC)과 같아야 한다.
        total = sum(lib.carbon_total(lib.records_for_kind(r.kind, ENV)[r.species], r.diameter, r.quantity, r.var2)
                    for r in p1.rows)
        self.assertAlmostEqual(round(total, 2), 198.89)

    def test_reader_accepts_english_headers_scientific_names_and_reports_bad_rows(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Regions"
        ws.append(["Site", "Width (m)", "Height (m)", "Site category"])
        ws.append(["A", 10, 30, "Post-fire site, natural restoration"])
        sp = wb.create_sheet("Species")
        sp.append(["Site", "Type", "Species", "Diameter (cm)", "Count",
                   "Initial survival (%)", "Annual survival (%)"])
        sp.append(["A", "Tree", "Pinus densiflora", 5, 10, 90, 98])
        sp.append(["A", "", "산철쭉", 1.2, 4, None, None])        # 구분 생략 → 라이브러리로 판정
        sp.append(["A", "Tree", "없는나무", 3, 1, 100, 100])
        sp.append(["A", "Shrub", "병꽃나무", 1.5, 2.5, 100, 100])  # 개체수 소수
        sp.append(["A", "Shrub", "병꽃나무", 1.5, 2, 120, 100])    # 생존율 범위 밖
        with TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "en.xlsx")
            wb.save(path)
            result = preset_io.read_preset(path, _names(lib.KIND_TREE), _names(lib.KIND_SHRUB),
                                           RESTORATION_ENVIRONMENTS)
        self.assertEqual(len(result.regions), 1)
        region = result.regions[0]
        self.assertEqual((region.name, region.width_m, region.height_m, region.environment), ("A", 10, 30, ENV))
        self.assertEqual([(r.kind, r.species) for r in region.rows], [("tree", "소나무"), ("shrub", "산철쭉")])
        self.assertEqual((region.rows[0].survival_initial, region.rows[0].survival_annual), (90.0, 98.0))
        self.assertEqual(len(result.errors), 3)

    def test_species_sheet_without_region_sheet_creates_default_site(self) -> None:
        wb = Workbook()
        ws = wb.active
        ws.title = "수종"
        ws.append(["구분", "수종", "직경(cm)", "개체수"])
        ws.append(["교목", "소나무", 5, 30])
        with TemporaryDirectory() as tmp:
            path = str(Path(tmp) / "one.xlsx")
            wb.save(path)
            result = preset_io.read_preset(path, _names(lib.KIND_TREE), _names(lib.KIND_SHRUB),
                                           RESTORATION_ENVIRONMENTS)
        self.assertEqual(result.errors, [])
        self.assertEqual(len(result.regions), 1)
        self.assertEqual((result.regions[0].width_m, result.regions[0].height_m), (20, 20))


class SurvivalExportTests(unittest.TestCase):
    def setUp(self) -> None:
        i18n.set_language(i18n.LANG_KO)

    @staticmethod
    def _category(all_vals, surv_vals, rows=()):
        years = list(range(len(all_vals)))
        return {
            "unit": "cm", "total": all_vals[0], "total_qty": 10, "rows": list(rows),
            "projection": {"years": years, "series": [{"label": "x", "checked": True, "carbon": all_vals}],
                           "total": None},
            "contribution": [],
            "survival": {"years": years, "all": all_vals, "with_survival": surv_vals,
                         "individuals": 10.0, "individuals_survival": [10.0 * s / a for s, a in zip(surv_vals, all_vals)]},
        }

    def test_single_and_combined_exports_include_survival_sheet(self) -> None:
        row = {"species": "소나무", "diameter": 5.0, "quantity": 10, "carbon": 1.0, "checked": True,
               "survival_initial": 90.0, "survival_annual": 98.0}
        payload = {
            "generated_at": "", "grand_total": 3.0,
            "region": {"name": "A", "environment": ENV, "area_w": 20, "area_h": 20, "area": 400},
            "tree": self._category([2.0, 4.0], [1.8, 3.5], [row]),
            "shrub": self._category([1.0, 2.0], [1.0, 2.0]),
            "images": [],
        }
        with TemporaryDirectory() as tmp:
            single = str(Path(tmp) / "single.xlsx")
            export_carbon1_to_excel(single, payload)
            ws = load_workbook(single, data_only=True)[tr("생존률_반영")]
            self.assertEqual([ws.cell(6, c).value for c in range(1, 8)], [1, 4.0, 3.5, 2.0, 2.0, 6.0, 5.5])

            combined = str(Path(tmp) / "combined.xlsx")
            export_all_regions_to_excel(combined, [payload], [])
            ws = load_workbook(combined, data_only=True)[tr("생존률_반영")]
            self.assertEqual([ws.cell(5, c).value for c in range(2, 5)], [1.8, 1.0, 2.8])


if __name__ == "__main__":
    unittest.main()
